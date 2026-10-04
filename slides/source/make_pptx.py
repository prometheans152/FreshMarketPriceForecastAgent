# -*- coding: utf-8 -*-
"""P6：產生正式可編輯 PPTX（15 主講頁＋4 附錄頁，16:9，微軟正黑體）。

執行：.venv/Scripts/python.exe slides/source/make_pptx.py
產出：slides/TW01_菜價預測_研究提案與初步成果_v1.pptx
數字來源：outputs/evaluation/public/public_test_results.json（凍結）＋ claims_evidence.csv。
"""
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
ASSETS = ROOT / "slides" / "assets"
OUT = ROOT / "slides" / "TW01_菜價預測_研究提案與初步成果_v1.pptx"

EVAL = json.loads((ROOT / "outputs/evaluation/public/public_test_results.json").read_text(encoding="utf-8"))

DARK = RGBColor(0x1F, 0x3B, 0x5C)
ACCENT = RGBColor(0x2E, 0x75, 0xB6)
GRAY = RGBColor(0x59, 0x59, 0x59)
LIGHT = RGBColor(0xF2, 0xF6, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xB0, 0x30, 0x30)
GREEN = RGBColor(0x2F, 0x7D, 0x32)

FONT = "微軟正黑體"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

SW, SH = prs.slide_width, prs.slide_height


def add_slide():
    return prs.slides.add_slide(BLANK)


def txbox(slide, l, t, w, h, text, size=18, color=DARK, bold=False, align=PP_ALIGN.LEFT, wrap=True):
    box = slide.shapes.add_textbox(Emu(int(l)), Emu(int(t)), Emu(int(w)), Emu(int(h)))
    tf = box.text_frame
    tf.word_wrap = wrap
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = ln
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.name = FONT
    return box


def title_bar(slide, title, subtitle=None):
    bar = slide.shapes.add_shape(1, 0, 0, SW, Inches(0.95))
    bar.fill.solid()
    bar.fill.fore_color.rgb = DARK
    bar.line.fill.background()
    tf = bar.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = title
    r.font.size = Pt(28); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT
    if subtitle:
        p2 = tf.add_paragraph()
        r2 = p2.add_run(); r2.text = subtitle
        r2.font.size = Pt(13); r2.font.color.rgb = RGBColor(0xC9, 0xD8, 0xE8); r2.font.name = FONT


def footer(slide, page, src):
    txbox(slide, Inches(0.4), Inches(7.05), Inches(9.5), Inches(0.4), f"來源：{src}", size=10, color=GRAY)
    txbox(slide, Inches(12.2), Inches(7.05), Inches(0.9), Inches(0.4), str(page), size=12, color=GRAY, align=PP_ALIGN.RIGHT)


def bullets(slide, items, l=Inches(0.55), t=Inches(1.25), w=Inches(12.3), size=18, gap=Pt(8)):
    box = slide.shapes.add_textbox(l, t, w, SH - t - Inches(0.6))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for it in items:
        text, lv = (it if isinstance(it, tuple) else (it, 0))
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = gap
        prefix = "▪ " if lv == 0 else "– "
        r = p.add_run(); r.text = prefix + text
        r.font.size = Pt(size - lv * 2)
        r.font.color.rgb = DARK if lv == 0 else GRAY
        r.font.name = FONT
    return box


def add_table(slide, rows, cols, l, t, w, h, data, header_fill=DARK, size=12):
    tblshape = slide.shapes.add_table(rows, cols, l, t, w, h)
    tbl = tblshape.table
    for ri, row in enumerate(data):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(size if ri else size)
                    r.font.name = FONT
                    r.font.color.rgb = WHITE if ri == 0 else DARK
                    r.font.bold = ri == 0
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if ri == 0 else (LIGHT if ri % 2 == 0 else WHITE)
    return tbl


def pic(slide, name, l, t, w=None, h=None):
    p = ASSETS / f"{name}.png"
    if w is not None:
        slide.shapes.add_picture(str(p), l, t, width=w)
    elif h is not None:
        slide.shapes.add_picture(str(p), l, t, height=h)
    else:
        slide.shapes.add_picture(str(p), l, t, w, h)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def badge(slide, l, t, text, color):
    sh = slide.shapes.add_shape(5, l, t, Inches(1.7), Inches(0.38))
    sh.fill.solid(); sh.fill.fore_color.rgb = color; sh.line.fill.background()
    tf = sh.text_frame; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text; r.font.size = Pt(12); r.font.bold = True
    r.font.color.rgb = WHITE; r.font.name = FONT


# ================= 主講 15 頁 =================

# --- 1 封面/摘要/狀態 ---
s = add_slide()
bg = s.shapes.add_shape(1, 0, 0, SW, SH); bg.fill.solid(); bg.fill.fore_color.rgb = DARK; bg.line.fill.background()
txbox(s, Inches(0.8), Inches(1.5), Inches(11.7), Inches(1.6), "現場菜價 7／14 天預測與分析代理人", size=40, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
txbox(s, Inches(0.8), Inches(3.1), Inches(11.7), Inches(0.6), "TW01・研究提案與初步成果（ midterm proposal + preliminary results）", size=20, color=RGBColor(0xC9, 0xD8, 0xE8), align=PP_ALIGN.CENTER)
txbox(s, Inches(0.8), Inches(4.1), Inches(11.7), Inches(1.5),
      "摘要：消費者在蔬果行輸入今天單品價格（上週同店價選填），系統以臺中公開市場均價模型\n預測 7／14 天方向並以現價錨定估計店價；A/B 輸入增益以真店到期價驗證（待驗證）。",
      size=15, color=WHITE, align=PP_ALIGN.CENTER)
txbox(s, Inches(0.8), Inches(6.2), Inches(11.7), Inches(0.5), "組員：待填（≤4人）｜日期：2026-10-04｜課程：AIoT 應用與資料分析", size=14, color=RGBColor(0x9F, 0xB8, 0xD0), align=PP_ALIGN.CENTER)
notes(s, "開場30秒：我們做的是『站在蔬果行，輸入今天標價，告訴你7天、14天後會不會更貴』。這是提案加初步成果：公開模型已封存測試，同店增益是明列的待驗證問題。（約40秒）")

# --- 2 情境與動機 ---
s = add_slide()
title_bar(s, "為什麼值得做：菜價會動，但消費者看不到未來", "動機（證據均可查）")
pic(s, "volatility", Inches(0.5), Inches(1.15), w=Inches(6.6))
bullets(s, [
    ("臺中 13 市場逐日均價開放資料存在（data.gov.tw 84539），但官方消費端工具只給「當日」行情", 0),
    ("官方研究：颱風後菜價第 1 週可漲 10–40%（農糧署《農政與農情》219期）", 0),
    ("本案資料實算（2025-10～2026-10，n=10,952）：7 天變動逾 ±5% 佔 10.6%；14 天 16.7%", 0),
    ("極端例：合作市場甘藍 2026-07-02 45→07-09 85 元/台斤（7 天 +88.9%）", 1),
    ("同時誠實：7 天完全同價佔 86.7–92.8% → 「多數不變、少數大變」，所以必須跟強基線比", 1),
], l=Inches(7.3), t=Inches(1.3), w=Inches(5.6), size=15)
footer(s, 2, "data/processed/features_labels.parquet 實算；農糧署219期；data.gov.tw/84539")
notes(s, "重點：不是『AI很熱門』，是菜價真的會動、且消費者工具只看當日。右下誠實揭露多數不變——這決定了我們後面一定要放不變價基線。極端例給聽眾具體感。（約60秒）")

# --- 3 問題與 RQ ---
s = add_slide()
title_bar(s, "研究問題", "RQ1–RQ3")
bullets(s, [
    ("RQ1｜公開市場固定期限方向預測：臺中均價的 7/14 天方向預測，相對「不變價」與「動量」兩個強基線表現如何？（macro-F1）", 0),
    ("RQ2｜現價／上週價增益：A（今天店價錨定）與 B（＋選填上週同店價）的增益，應如何以真實同店 7/14 天到期價驗證？——目前零證據，屬待驗證", 0),
    ("RQ3｜低負擔與可追溯：單品兩欄輸入、附來源與驗證狀態的呈現，能否兼顧可用與可核對？（人因評估為未來工作）", 0),
    ("使用情境：高麗菜今天 32 元/台斤 → 7 天後（10-11）、14 天後（10-18）相對今天是漲、平還是跌？", 1),
], size=18)
rq_cards = [
    ("RQ1 → C1", "固定期限評估基準", "已驗證", GREEN),
    ("RQ2 → C2", "A/B 輸入增益", "待驗證", RGBColor(0xC0, 0x7A, 0x00)),
    ("RQ3 → C3", "低負擔可追溯呈現", "初步", GRAY),
]
for i, (top_t, mid_t, st, c) in enumerate(rq_cards):
    x = Inches(0.55 + i * 4.25)
    card = s.shapes.add_shape(5, x, Inches(4.6), Inches(4.0), Inches(1.9))
    card.fill.solid(); card.fill.fore_color.rgb = LIGHT; card.line.color.rgb = c; card.line.width = Pt(1.5)
    tf = card.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = top_t; r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = DARK; r.font.name = FONT
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = mid_t; r2.font.size = Pt(15); r2.font.color.rgb = DARK; r2.font.name = FONT
    p3 = tf.add_paragraph(); p3.alignment = PP_ALIGN.CENTER
    r3 = p3.add_run(); r3.text = f"證據狀態：{st}"; r3.font.size = Pt(13); r3.font.bold = True; r3.font.color.rgb = c; r3.font.name = FONT
footer(s, 3, "docs/RESEARCH_MOTIVATION.md（RQ 定義與證據）；狀態對應 docs/CONTRIBUTIONS.md")
notes(s, "三個RQ分別對應三貢獻。強調RQ2明講零證據——我們不假裝已驗證。（約45秒）")

# --- 4 前作分類 ---
s = add_slide()
title_bar(s, "相關研究：三類前作", "分類討論（詳矩陣與引用清單）")
add_table(s, 5, 4, Inches(0.5), Inches(1.2), Inches(12.3), Inches(4.6), [
    ["類別", "代表前作", "已做到", "未做／未知"],
    ["批發・市場價格預測", "陳泯宏2024[R01]；陳澤維[R09]；林沂臻2017[R03]；NCHU 2025[R07]", "批發/市場均價預測（ARIMA/LSTM/HHT/氣候探勘）", "消費者店價輸入；固定7/14日曆期限；同店到期"],
    ["零售・私人資料", "農業大數據平台2018[R04]；李岱鑫2015[R02]", "電子發票零售估計預測；私人vs公開資料誤差比較", "消費者輸入；單店真值到期驗證"],
    ["消費者行情工具", "白菜即時[R05]；農產分析[R06]", "AI走勢預測（不透明）；輸入自己買的價格比對當日", "未來期限預測[R06]；店價輸入與指標公開[R05]"],
    ["官方動向", "農糧署115年度徵件[R10]", "生成式AI＋預測ML＋農產價格獲政策重視", "生產端決策；非消費者情境"],
], size=13)
footer(s, 4, "docs/RELATED_WORK_MATRIX.md（10件；UNKNOWN 12處如實標）")
notes(s, "過半篇幅直接談菜價。三類各自『已做到』講清楚——特別是農產分析已讓人輸入自己買的價格、大數據平台已做零售預測，我們的差異在組合而非元件。（約50秒）")

# --- 5 前作比較與差異 ---
s = add_slide()
title_bar(s, "與最接近前作的差異", "查核範圍內未見相同組合；不宣稱全球首次")
add_table(s, 5, 5, Inches(0.4), Inches(1.2), Inches(12.5), Inches(4.2), [
    ["維度", "李岱鑫2015 [R02]", "大數據平台2018 [R04]", "農產分析 [R06]", "本案 TW01"],
    ["輸入", "中盤商私人紀錄", "電子發票（無輸入）", "使用者買價（比對當日）", "今天店價＋選填上週同店價"],
    ["目標", "下一週批發價", "零售通路段估計", "當日價差合理性", "從今天起 t+7/t+14 方向"],
    ["真值驗證", "有（批發誤差比較）", "網路查價對照", "無（當日）", "同店到期價（設計中，T16）"],
    ["主要限制", "批發；通路資料", "均價估計非店價", "不預測未來", "B 增益零證據；僅一年臺中"],
], size=13)
txbox(s, Inches(0.5), Inches(5.6), Inches(12.3), Inches(0.9),
      "三組缺口（詳 RESEARCH_GAP.md）：①今天店價基準＋固定日曆期限的評估基準 ②消費者選填上週價＋同店到期配對驗證 ③低輸入負擔＋主張狀態可追溯的介面", size=14, color=GRAY)
footer(s, 5, "docs/RESEARCH_GAP.md；未重現前作協議，不跨資料比分數")
notes(s, "差異是組合不是元件。強調『查核範圍內未見』的用語紀律，以及不跨資料排名。（約45秒）")

# --- 6 三貢獻 ---
s = add_slide()
title_bar(s, "三項貢獻與證據狀態", "如實標示：已驗證／待驗證／初步")
cards = [
    ("C1（已驗證）", "固定 7/14 日曆期限的臺中零售\n方向預測評估基準：強基線、\n時間切分、封存測試，含負向結果", GREEN),
    ("C2（待驗證）", "現價／上週價 A/B 輸入與真店\n到期配對驗證協議；B 目前為\nα=0.3 固定實驗規則（未訓練）", RGBColor(0xC0, 0x7A, 0x00)),
    ("C3（初步）", "主張狀態可追溯的消費者端流程\n與原型：來源、期限、驗證範圍\n標示貫穿文件與介面", GRAY),
]
for i, (h, body, c) in enumerate(cards):
    x = Inches(0.5 + i * 4.25)
    card = s.shapes.add_shape(5, x, Inches(1.3), Inches(4.0), Inches(0.55))
    card.fill.solid(); card.fill.fore_color.rgb = c; card.line.fill.background()
    p = card.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = h; r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT
    txbox(s, x, Inches(2.0), Inches(4.0), Inches(2.6), body, size=14, color=DARK)
bullets(s, [
    ("C1 已有可重現證據：方向 macro-F1 0.4218/0.4169；勝全持平（0.3092/0.2954）、未勝動量（0.4226/0.4306）；點 MAE 未勝不變價", 0),
    ("C2 的增益是本提案明列的待驗證問題；C3 的使用者研究未執行——兩者皆不寫成已證實", 1),
], t=Inches(4.9), size=15)
footer(s, 6, "docs/CONTRIBUTIONS.md；outputs/evaluation/public/")
notes(s, "顏色即狀態：綠=已驗證、黃=待驗證、灰=初步。誠實是這頁的重點。（約45秒）")

# --- 7 資料 ---
s = add_slide()
title_bar(s, "資料：臺中公有市場每日零售均價", "資料卡與品質報告")
bullets(s, [
    ("範圍：2025-10-04 ～ 2026-10-02；14 市場（排除「(停)建國」沿用價密集 → 13）；364 觀測日無斷點", 0),
    ("單位：元/台斤（1 台斤 = 0.6 kg）；0 = 當日無價 → 缺測（905 格），不是 0 元", 0),
    ("首版品項：甘藍（平地高麗菜）、青江白菜、絲瓜；無個店 ID——是市場均價，不是某家店", 0),
    ("配對：精確日曆日 t+7/t+14，兩端正價才成標籤；候選標籤 7 天 3,615 列／14 天 3,523 列（2026-10-03 實算）", 1),
    ("限制：僅一年、單一城市；發布延遲以保守假設 +1 天處理（未實測）", 1),
    ("來源：data.gov.tw 84539（政府開放授權1.0）；快照 SHA256 172dffc7…（可稽核）", 1),
], size=15)
stats = [("364", "觀測日無斷點"), ("13", "納入市場"), ("3", "首版品項"), ("905", "缺測格（0=無價）")]
for i, (num, lab) in enumerate(stats):
    x = Inches(0.55 + i * 3.15)
    card = s.shapes.add_shape(5, x, Inches(4.75), Inches(2.9), Inches(1.7))
    card.fill.solid(); card.fill.fore_color.rgb = LIGHT; card.line.color.rgb = ACCENT; card.line.width = Pt(1.25)
    tf = card.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = num
    r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = ACCENT; r.font.name = FONT
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = lab
    r2.font.size = Pt(13); r2.font.color.rgb = DARK; r2.font.name = FONT
footer(s, 7, "docs/DATA_CARD.md；docs/DATA_QUALITY_REPORT.md")
notes(s, "資料小但乾淨：無斷點、鍵唯一。強調 0=無價 與 均價≠店價 這兩個誠實邊界，後面方法都由此展開。（約40秒）")

# --- 8 架構 ---
s = add_slide()
title_bar(s, "系統架構：現況與暫緩", "實線=已實作並本機驗證；虛線=暫緩（未部署）")
shapes_def = [
    (Inches(0.6), Inches(1.6), Inches(3.0), Inches(1.5), "手機網頁\n輸入/A・B並列/保存", LIGHT, True),
    (Inches(4.2), Inches(1.6), Inches(3.0), Inches(1.5), "HTTP API\n/forecast /agent /records", LIGHT, True),
    (Inches(7.8), Inches(1.6), Inches(3.0), Inches(1.5), "ForecastEngine\nRF(7天)/HGB(14天) 封存v0.1.0", LIGHT, True),
    (Inches(7.8), Inches(3.6), Inches(3.0), Inches(1.2), "公開歷史\npublic_history（凍結快照）", LIGHT, True),
    (Inches(4.2), Inches(3.6), Inches(3.0), Inches(1.2), "規則代理人\n（零LLM，展示用）", LIGHT, True),
    (Inches(4.2), Inches(5.3), Inches(3.0), Inches(1.1), "GLM API 後端\n金鑰未配置", RGBColor(0xFA, 0xF3, 0xE3), False),
    (Inches(7.8), Inches(5.3), Inches(3.0), Inches(1.1), "Qwen3.8-27B 端點\nSERVER 未部署", RGBColor(0xFA, 0xF3, 0xE3), False),
    (Inches(0.6), Inches(5.3), Inches(3.0), Inches(1.1), "Hermes 0.19.0 框架\n已離線驗證未接真模型", RGBColor(0xFA, 0xF3, 0xE3), False),
]
for l, t, w, h, text, fill, solid in shapes_def:
    sh = s.shapes.add_shape(5, l, t, w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = DARK if solid else GRAY
    sh.line.dash_style = None if solid else 4  # dash
    sh.line.width = Pt(1.5) if solid else Pt(1)
    tf = sh.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text; r.font.size = Pt(13); r.font.name = FONT
    r.font.color.rgb = DARK; r.font.bold = solid
txbox(s, Inches(0.6), Inches(6.6), Inches(12.2), Inches(0.5), "虛線區（黃底）＝D44 暫緩：代理人語言模型後端整合未接線，介面以規則模式作展示載體；不宣稱已通過課程全部要求", size=12, color=GRAY)
footer(s, 8, "design.md §1；ROUND5_REPAIR_REPORT.md（工程細節）")
notes(s, "一張圖講清楚哪些是真的（上排與歷史）、哪些暫緩（黃底虛線）。誠實邊界直接畫在圖上。（約40秒）")

# --- 9 時間切分 ---
s = add_slide()
title_bar(s, "7／14 天目標與時間切分", "防洩漏設計")
pic(s, "time_split", Inches(0.7), Inches(1.3), w=Inches(10.2))
bullets(s, [
    ("標籤契約（凍結於看測試前）：r_h = m(t+h)/m(t) − 1；|r| ≤ 5% 為持平；缺到期價＝無標籤，不是「未漲」", 0),
    ("同日 13 市場進同區（相關觀測不跨切分稀釋）；14+1 天成熟清除；封存測試一次性（清單 SHA256 4e136574…）", 0),
    ("期限永遠從使用者今天 t 起算，不因行情截止日較舊而改起點", 1),
], t=Inches(4.65), size=15)
footer(s, 9, "docs/LABEL_CONTRACT.md；manifest.json splits")
notes(s, "切分圖+三個防洩漏要點。成熟清除是聽眾最容易忽略的細節，點一句即可。（約40秒）")

# --- 10 A/B 方法與數字例 ---
s = add_slide()
title_bar(s, "A/B 兩種輸入與一個數字例", "A=現價錨定；B=上週價實驗規則")
txbox(s, Inches(0.5), Inches(1.2), Inches(6.2), Inches(0.5), "公式", size=17, bold=True)
txbox(s, Inches(0.5), Inches(1.7), Inches(6.2), Inches(2.6),
      "A 店價估計：p̂(t+h) = p(t) × (1 + r̂_h)\nB 調整：r_B = r_A + α × (Δ店價7日 − Δ公開同期)\n　α = 0.3（開發期固定；未經同店資料訓練）\n\n錨定不變向：p>0 時 A 的方向 ≡ r̂ 的方向\n→ 輸入現價只改變金額，不改善方向判斷",
      size=15)
txbox(s, Inches(0.5), Inches(4.4), Inches(6.2), Inches(0.5), "輸入負擔", size=17, bold=True)
txbox(s, Inches(0.5), Inches(4.9), Inches(6.2), Inches(1.4), "A：1 個必填（今天價格）＋品項\nB：再加 2 個選填（上週價＋實際日期）\n上週不可比（7±2天外/換店/換規格/跨品項）→ 自動回 A 並說明", size=15)
# 右側數字例卡片
card = s.shapes.add_shape(5, Inches(7.1), Inches(1.2), Inches(5.7), Inches(5.3))
card.fill.solid(); card.fill.fore_color.rgb = LIGHT; card.line.color.rgb = ACCENT
txbox(s, Inches(7.35), Inches(1.4), Inches(5.3), Inches(0.5), "數字例（甘藍，2026-10-04）", size=16, bold=True)
txbox(s, Inches(7.35), Inches(1.95), Inches(5.3), Inches(4.4),
      "輸入：今天 32 元/台斤；上週同店 28 元\n\n模型（市場）：r̂_7 = −4.1% → 持平\n　　　　　　r̂_14 = −0.9% → 持平\n\nA：7天後 ≈ 30.7 元、14天後 ≈ 31.7 元\n\nB：店價週變化 = 32/28−1 = +14.3%\n　公開同期 +1.0%（例）\n　r_B = −4.1% + 0.3×13.3% = −0.1%（仍持平）\n\n回應附：資料截止 10-02、點估計僅供參考、\nB 為實驗性規則未經同店驗證",
      size=14)
footer(s, 10, "fixtures（封存模型輸出）；store_adapter.py")
notes(s, "左公式右例子。錨定不變向是本題最重要的誠實推論：輸入現價不會讓方向變準——這正是B存在的理由，而B的增益必須等真店到期驗證。數字例一步一步念。（約70秒）")

# --- 11 模型與比較設計 ---
s = add_slide()
title_bar(s, "模型與公平比較設計", "選型凍結於看測試前")
bullets(s, [
    ("候選：Ridge / HistGradientBoosting / RandomForest / 分類樹；以驗證集 macro-F1 選型（凍結後才碰封存測試）", 0),
    ("選定：7 天 RF(300樹)、14 天 HGB(200輪)；seed 42；特徵只用 t−1 前可見資料（價齡/滾動統計/報酬…12 維，無市場 one-hot）", 0),
    ("兩個強基線（與模型同測試列）：不變價（全持平）、動量（過去同長度變化延續）", 0),
    ("指標：方向 macro-F1（持平多數使 accuracy 失真）；r 的 MAE（非元、非準確率）", 0),
    ("未做（列為計畫，不補造）：參數敏感度、特徵消融、與近期代表方法在同資料上的直接比較", 1),
], size=16)
footer(s, 11, "docs/MODEL_SELECTION.md；training_results.json")
notes(s, "公平比較的重點是基線強、指標抗不平衡、選型時序凍結。未做的事直接列。（約35秒）")

# --- 12 結果 ---
s = add_slide()
title_bar(s, "封存測試結果（公開市場均價）", "正向與負向並列——這是誠實的成績單")
pic(s, "f1_vs_baselines", Inches(0.45), Inches(1.1), w=Inches(5.7))
pic(s, "mae_negative", Inches(0.45), Inches(4.21), w=Inches(5.7))
bullets(s, [
    ("方向 macro-F1：7 天 0.4218、14 天 0.4169", 0),
    ("勝全持平基線（0.3092 / 0.2954）", 1),
    ("未勝動量基線（0.4226 / 0.4306）", 1),
    ("點估計 MAE 未勝不變價（0.0664 vs 0.0436；0.0983 vs 0.0748）→ 點估計僅供參考", 0),
    ("up 召回 ≈10%/9%：模型偏保守報持平，「會漲」警示價值受限", 0),
    ("無檢定 → 不寫「顯著」；樣本相關 → 不當獨立", 1),
], l=Inches(7.1), t=Inches(1.3), w=Inches(5.8), size=14)
footer(s, 12, "public_test_results.json（SHA256 2d28acd2…，可重算）；n=2,085/1,859")
notes(s, "本頁是提案的核心誠實：模型有訊號（勝持平）但不是最強（未勝動量）、點誤差輸給不變價。這些負向結果定義了後續工作。（約60秒）")

# --- 13 介面 ---
s = add_slide()
title_bar(s, "介面現況與限制", "原型截圖（2026-10-03，本機瀏覽器）")
shot = ROOT / "outputs/e2e/ui_ab_mobile_390x844.png"
if shot.exists():
    s.shapes.add_picture(str(shot), Inches(0.7), Inches(1.2), height=Inches(5.4))
bullets(s, [
    ("已實作（本機）：表單/自然語言兩入口、A/B 並列、保存→到期核對、匯出/刪除（request_id 綁定伺服器預測）", 0),
    ("主張狀態進畫面：資料截止日、驗證範圍、B 實驗性警示", 0),
    ("限制：本機原型（rule 模式，零 LLM）；未部署外網（T32 HOLD）；代理人後端整合暫緩", 1),
    ("截圖為 2026-10-03 原型；與稿件不一致時以文件為準", 1),
], l=Inches(4.3), t=Inches(1.5), w=Inches(8.5), size=15)
footer(s, 13, "outputs/e2e/ui_ab_mobile_390x844.png；docs/UI_FLOW.md")
notes(s, "一張真截圖勝過描述。右側把限制講清楚：rule模式、未外網、後端暫緩。（約35秒）")

# --- 14 驗證計畫/OpenSpec ---
s = add_slide()
title_bar(s, "分階段驗證計畫與規格管理", "design.md 藍圖＋OpenSpec 變更管理")
add_table(s, 4, 3, Inches(0.5), Inches(1.25), Inches(12.3), Inches(3.0), [
    ["C/RQ", "證據", "狀態"],
    ["C1 / RQ1", "封存測試 JSON+CSV（可重算）；基線並列", "已驗證（含負向）"],
    ["C2 / RQ2", "A/B 協議＋凍結預測保存；真店 7/14 天到期價", "協議就緒／增益待驗證（T16）"],
    ["C3 / RQ3", "原型＋主張→來源→狀態追溯鏈（claims_evidence.csv）", "初步／人因未測"],
], size=14)
bullets(s, [
    ("規格管理：OpenSpec 1.14.0（局部安裝）；change「prepare-tw01-proposal-package」4/4 artifacts、validate --strict 通過", 0),
    ("時間表（待填日期）：真店收集啟動 → 參數敏感度 → SERVER 部署與外網 → 期末全案驗收（各項單獨驗收，不互相冒充）", 0),
], t=Inches(4.6), size=15)
footer(s, 14, "openspec_cli_evidence.txt；docs/OPENSPEC_WORKFLOW.md")
notes(s, "把驗證狀態當作交付的一部分來管理。OpenSpec 是真的有初始化與validate證據的。（約35秒）")

# --- 15 結論 ---
s = add_slide()
title_bar(s, "結論與下一步")
bullets(s, [
    ("回答 RQ1：固定期限＋強基線的零售評估可行且可重現；模型方向訊號存在（勝持平）但未勝動量、點誤差未勝不變價", 0),
    ("回答 RQ2/RQ3：協議與原型就緒；增益與人因皆待驗證——這是接下來的工作，不是已完成的成果", 0),
    ("下一步：①啟動真店到期收集（T16） ②參數敏感度與近期方法同資料比較 ③SERVER 部署與外網評估（T24/T29/T32） ④使用者研究", 0),
    ("一句話總結：我們交付的不是「很準的預測」，而是一個誠實可驗證的預測與增益評估框架", 1),
], size=17)
txbox(s, Inches(0.55), Inches(5.6), Inches(12.2), Inches(0.8), "感謝聆聽——歡迎提問（附錄：分品項/混淆矩陣/文獻/公式）", size=16, bold=True, align=PP_ALIGN.CENTER)
footer(s, 15, "docs/PROPOSAL_REPORT.md W–Z")
notes(s, "收尾30秒：強調貢獻是框架與誠實，不是神奇準確率。指向附錄應付提問。（約35秒）")

# ================= 附錄 =================

# A1 分品項
s = add_slide()
title_bar(s, "附錄 A1：分品項 macro-F1", "公開市場均價測試")
pic(s, "by_crop", Inches(1.2), Inches(1.3), w=Inches(9.0))
txbox(s, Inches(1.2), Inches(5.8), Inches(10.5), Inches(0.8), "甘藍方向事件最少（7天 |r|>5% 僅 6.8%），macro-F1 最低屬預期；絲瓜事件多、訊號較強。樣本數見圖上標註。", size=14, color=GRAY)
footer(s, "A1", "public_test_results.json by_crop")
notes(s, "備詢用：品項差異來自事件頻率不同。")

# A2 混淆矩陣
s = add_slide()
title_bar(s, "附錄 A2：7 天混淆矩陣", "列=實際；RF")
pic(s, "confusion_h7", Inches(0.6), Inches(1.3), w=Inches(5.3))
bullets(s, [
    ("up 召回 = 16/158 ≈ 10.1%：預測過保守", 0),
    ("down 召回 = 57/124 ≈ 46%：跌勢抓得相對好", 0),
    ("flat 精確率 0.90：報持平時多半正確", 0),
    ("門檻 ±5% 的敏感度分析未執行（計畫）", 1),
], l=Inches(6.35), t=Inches(1.8), w=Inches(6.2), size=15)
footer(s, "A2", "public_test_results.json confusion")
notes(s, "備詢用：為何macro-F1輸動量——up召回低是主因之一。")

# A3 文獻
s = add_slide()
title_bar(s, "附錄 A3：引用清單（[R01…R10]）")
refs = [
    "R01 陳泯宏(2024) 國產大宗蔬菜價格預測模型之比較. 臺大碩士. (摘要)",
    "R02 李岱鑫(2015) 高麗菜批發價格預測(公開+私人). 清華碩士. (摘要)",
    "R03 林沂臻(2017) HHT菜價預測. 中山醫大碩士. (摘要)",
    "R04 農委會×財政資訊中心(2018) 農業大數據平台零售估計預測. (報導全文)",
    "R05 flux6lab. 白菜即時 App. (商店文案, 2026-08)",
    "R06 鹿中創客. 農產分析 App. (商店文案, 2026-07)",
    "R07 YuehChun(2025) NCHU-Data-Mining-Project. GitHub. (repo)",
    "R08 洪士育(n.d.) 氣候→蔬菜價格. 北科大. (摘要片段)",
    "R09 陳澤維(約2017) 臺灣蔬果批發價格預測. 政大碩士. (摘要)",
    "R10 農糧署(2026) 115年度AI+ML農產價格先導研究徵件. (PDF全文)",
]
bullets(s, [(r, 0) for r in refs], l=Inches(0.6), t=Inches(1.15), w=Inches(12.2), size=13, gap=Pt(4))
footer(s, "A3", "docs/REFERENCES.md（完整書目與連結）")
notes(s, "完整書目與讀取層級在REFERENCES.md。")

# A4 公式與需求對照
s = add_slide()
title_bar(s, "附錄 A4：公式彙整與課程要求對照")
txbox(s, Inches(0.5), Inches(1.15), Inches(6.1), Inches(5.4),
      "公式彙整\n\nr_h = m(t+h)/m(t) − 1\n方向：|r|≤5% flat；> +5% up；< −5% down\n\nA：p̂(t+h) = p(t)·(1+r̂_h)\nB：r_B = r_A + 0.3·(Δ店價7日 − Δ公開同期)\n單位：1 台斤 = 0.6 kg；50元/kg ≡ 30元/台斤\n\n錨定不變向（代數）：p>0 ⇒ sign(p·(1+r̂)−p)=sign(r̂)\n適用條件：乘法錨定；B 之加法項可改變方向（故待驗證）", size=15)
add_table(s, 6, 2, Inches(6.9), Inches(1.2), Inches(5.9), Inches(3.6), [
    ["課程要求", "本案對應"],
    ["一定要有模型", "RF/HGB 封存評估（T15）"],
    ["希望使用 Agent", "介面原型＋後端整合暫緩（如實標示）"],
    ["量化評估", "macro-F1/MAE/混淆矩陣＋基線"],
    ["網路上可操作", "本機原型；外網部署待 SERVER（T32）"],
    ["design.md 藍圖", "本次重整＋OpenSpec 管理"],
], size=12)
footer(s, "A4", "design.md §3；逐字稿行1710/1711/2832（歷史快照）")
notes(s, "備詢用：公式一頁看懂；課程要求對照也誠實標Agent暫緩。")

prs.save(str(OUT))
print("saved:", OUT)
print("slides:", len(prs.slides.__iter__.__self__._sldIdLst))
