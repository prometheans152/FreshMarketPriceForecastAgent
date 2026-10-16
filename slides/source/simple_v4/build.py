# -*- coding: utf-8 -*-
"""TW01 v4: white-background research presentation. No model or product changes."""
from __future__ import annotations
import argparse
from datetime import datetime
import hashlib
import json
from pathlib import Path
import zipfile
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

VERSION='v4'
NAME='菜價預測_研究提案_簡潔版_v4'
FONT='Microsoft JhengHei'
INK='222B33'; ACCENT='315A70'; MUTED='606B74'; LINE='CDD4DA'; LIGHT='F1F4F6'; WHITE='FFFFFF'
WIDTH,HEIGHT=13.333333,7.5
MAIN_COUNT=16
SOURCES={
 'D1':'臺中市政府公有零售市場每日蔬果價格表，data.gov.tw資料集84539；https://data.gov.tw/dataset/84539 。使用既有2026-10-03快照，不重新下載。',
 'D2':'outputs/proposal_delivery/run3_20261004_01/stats_master.json: descriptive；data/processed/features_labels.parquet的既存direction。含各時間區的歷史市場×品項×起算日配對；排除insufficient_data，正式±5%邊界含eps=1e-9。',
 'E1':'outputs/evaluation/public/public_test_results.json: horizons.h=7 / horizons.h=14；SHA256 2d28acd2f2c3e5ce566b9095c6a85477db8faf858d9ce29e974caaedc25674b8。固定測試列，未重算模型輸出、未重訓。',
 'E2':'outputs/evaluation/public/public_test_results.json: by_crop；同一封存測試的分品項描述。',
 'M1':'docs/LABEL_CONTRACT.md；src/aiot_price/data/labels.py；docs/PROPOSAL_REPORT.md。h=7/14個日曆日；t為本次可信報價基準日；缺到期正價格者無標籤。',
 'M2':'src/aiot_price/models/store_adapter.py: ALPHA_EXPERIMENTAL=0.3；docs/STORE_DATA_PROTOCOL.md；docs/PROPOSAL_REPORT.md。A為乘法錨定；B為固定規則，不是學得的增益。',
 'M3':'outputs/training/training_results.json；models/releases/0.1.0/manifest.json；src/aiot_price/models/train.py。候選Ridge、RFRegressor、HGBRegressor、DummyMajority、HGBClassifier；12個數值特徵＋crop_id one-hot，無市場one-hot。',
 'S1':'outputs/proposal_delivery/run3_20261004_01/stats_master.json: splits, purge.segments；outputs/splits/split_list.csv。成熟日=t+h+1。图示日期为保留樣本涵蓋期，不表示每一天每個期限都保留。',
 'A1':'src/aiot_price/api/app.py: /forecast, /agent/query, /records；src/aiot_price/models/predict.py: ForecastEngine；src/aiot_price/api/prediction_records.py: private/predictions。依既有原型整理的功能關係圖，不是完整系統驗收宣稱。',
 'U1':'outputs/e2e/ui_ab_mobile_390x844.png，2026-10-03本機原型截圖。本版只裁出上方7天A/B結果區並明示裁切，未改內部數字、未偽造新查詢。',
 'P1':'docs/PROPOSAL_REPORT.md: RQ1–RQ3、C1–C3及後續工作；docs/CONTRIBUTIONS.md；TASKS.md。使用者2026-10-04指示產品代理人整合與SERVER暫緩。',
 'R01':'陳泯宏（2024），《國產大宗蔬菜價格預測模型之比較》，臺灣大學統計碩士學位學程。DOI 10.6342/NTU202400875；https://tdr.lib.ntu.edu.tw/handle/123456789/92775 。沿用既有摘要查核，未讀全文。',
 'R02':'李岱鑫（2015），《高雄果菜運銷公司高麗菜批發價格預測之研究—蒐集私有資料的價值》，清華大學。https://etd.lib.nycu.edu.tw/cgi-bin/gs32/hugsweb.cgi?o=dnthucdr&s=id%3D%22GH02102077503%22.&searchmode=basic 。沿用書目與摘要查核；實際評估指標仍UNKNOWN。',
 'R05':'flux6lab，白菜即時官方頁 https://cabbagelive.com/ 。沿用已讀產品文案：趨勢預測／零售換算。未安裝實測，不把文案當成效果證據，未載明功能保持未知。',
 'R06':'鹿中創客，農產分析Google Play頁 https://play.google.com/store/apps/details?id=com.bvc_dev_team.androidapp&hl=zh-TW 。沿用已讀文案：輸入購買價、換算及當日行情比對；未安裝實測，其餘功能不作不存在的斷言。',
 'H17':'同課程Team17，《蔬果價格預測：LightGBM的最佳採購指南》（2025），本機歷屆/AIoT_歷屆期末簡報/regular_ii/Team17_蔬果價格預測_LightGBM採購指南.pdf，第7–12頁；原頁 https://app.notion.com/p/Team-17-2d3c4752981180e7985ed1361620f46a 。只依作者簡報，不是已獨立重現的研究結果。',
 'C1':'會議錄音 10.txt:535–573、787–899、1738–1764、2080–2097；教授L4-1 https://app.notion.com/p/AIoT-DA-L4-1-Proposal-1-3ebc4752981180d8842fd7e253962c4e 。沿用已讀正文與逐字稿；A–Z為內容對照，不是全部需放主畫面的逐字標題。',
 'STYLE':'本版參考：LSTM澆水第4/14/22頁的簡圖與實物展示；蜜蜂健康第5–14頁的背景/動機/前作論述；GUARD節錄第4/5張的總覽與局部解釋；Team17第7–10頁的資料與結果安排。不複製他組數據、圖片、成果或新功能。'
}

def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def color(v): return RGBColor.from_string(v)

def no_effects(shape):
    """Remove theme-inherited shadows as well as explicit effect lists."""
    sp=shape._element.spPr
    for tag in ['a:effectLst','a:effectDag']:
        old=sp.find(qn(tag))
        if old is not None:sp.remove(old)
    sp.append(OxmlElement('a:effectLst'))
    style=shape._element.find(qn('p:style'))
    if style is not None:
        ref=style.find(qn('a:effectRef'))
        if ref is not None:ref.set('idx','0')

class Deck:
    def __init__(self,root,out,evidence):
        self.root,self.out,self.evidence=root,out,evidence
        self.prs=Presentation();self.prs.slide_width=Inches(WIDTH);self.prs.slide_height=Inches(HEIGHT)
        cp=self.prs.core_properties;cp.title='TW01 現場菜價7／14天預測與分析代理人';cp.subject='研究提案與初步成果；簡潔研究版v4';cp.author='TW01 project'
        cp.created=cp.modified=datetime(2026,10,4)
        self.notes,self.layout=[],[]
        self.eval=json.loads((root/'outputs/evaluation/public/public_test_results.json').read_text(encoding='utf-8'))
        self.stats=json.loads((root/'outputs/proposal_delivery/run3_20261004_01/stats_master.json').read_text(encoding='utf-8'))
        assert sha(root/'outputs/evaluation/public/public_test_results.json')=='2d28acd2f2c3e5ce566b9095c6a85477db8faf858d9ce29e974caaedc25674b8','Frozen evaluation changed; inspect before authoring.'
        assert self.stats['descriptive']['h7|ALL']['nonflat']==1149
        assert self.stats['descriptive']['h14|ALL']['nonflat']==1773
        self.shot=root/'outputs/e2e/ui_ab_mobile_390x844.png'
        if not self.shot.is_file():raise FileNotFoundError('Required historical UI screenshot is missing.')

    def text(self,s,x,y,w,h,text,size=22,bold=False,tint=INK,align='left',linespace=1.16):
        box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=box.text_frame
        tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.vertical_anchor=MSO_ANCHOR.TOP
        for i,line in enumerate(str(text).split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
            p.line_spacing=linespace;p.space_after=Pt(3)
            r=p.add_run();r.text=line;r.font.name=FONT;r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=color(tint)
            ea=OxmlElement('a:ea');ea.set('typeface',FONT);r._r.get_or_add_rPr().append(ea)
        self.layout.append({'page':len(self.prs.slides),'kind':'text','x':x,'y':y,'w':w,'h':h,'size':size,'text':str(text)})
        return box

    def line(self,s,x1,y1,x2,y2,tint=LINE,width=1,arrow=False):
        obj=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));obj.line.color.rgb=color(tint);obj.line.width=Pt(width)
        no_effects(obj)
        if arrow:
            ln=obj.line._get_or_add_ln();end=OxmlElement('a:tailEnd');end.set('type','triangle');ln.append(end)
        return obj

    def box(self,s,x,y,w,h,fill=WHITE,stroke=LINE):
        sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
        sh.fill.solid();sh.fill.fore_color.rgb=color(fill);sh.line.color.rgb=color(stroke);sh.line.width=Pt(1)
        no_effects(sh)
        return sh

    def page(self,title,section,footer):
        s=self.prs.slides.add_slide(self.prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=color(WHITE)
        self.text(s,.65,.30,11.85,.25,section,11,tint=ACCENT)
        self.text(s,.65,.72,12,.61,title,29,bold=True)
        self.line(s,.65,1.47,12.68,1.47)
        self.text(s,.65,7.02,11.6,.23,footer,10,tint=MUTED,linespace=1)
        self.text(s,12.02,7,.66,.27,str(len(self.prs.slides)),11,tint=MUTED,align='right')
        return s

    def takeaway(self,s,text,y=6.34):
        self.line(s,.65,y-.17,1.15,y-.17,tint=ACCENT,width=2)
        self.text(s,.65,y,12,.52,text,20,bold=True,tint=ACCENT)

    def points(self,s,items,x=.75,y=1.9,w=11.8,size=22,gap=1.18):
        for i,(head,body) in enumerate(items):
            yy=y+i*gap
            self.text(s,x,yy,.48,.43,f'{i+1:02d}',16,bold=True,tint=ACCENT)
            self.text(s,x+.68,yy,w-.68,.46,head,size,bold=True)
            self.text(s,x+.68,yy+.52,w-.68,.70,body,size-2,tint=MUTED)

    def table(self,s,x,y,width,heights,colwidths,rows,size=19,highlight=()):
        sh=s.shapes.add_table(len(rows),len(colwidths),Inches(x),Inches(y),Inches(width),Inches(sum(heights)));tbl=sh.table
        for i,w in enumerate(colwidths):tbl.columns[i].width=Inches(w)
        for i,h in enumerate(heights):tbl.rows[i].height=Inches(h)
        for ri,row in enumerate(rows):
            for ci,txt in enumerate(row):
                c=tbl.cell(ri,ci);c.text=str(txt);c.margin_left=c.margin_right=Inches(.13);c.margin_top=c.margin_bottom=Inches(.09);c.vertical_anchor=MSO_ANCHOR.MIDDLE
                c.fill.solid();c.fill.fore_color.rgb=color(LIGHT if ri==0 or ri in highlight else WHITE)
                pr=c._tc.get_or_add_tcPr()
                for edge in ['lnL','lnR','lnT','lnB']:
                    found=pr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}'+edge)
                    if found is not None:pr.remove(found)
                    ln=OxmlElement('a:'+edge);ln.set('w','6350');sol=OxmlElement('a:solidFill');cc=OxmlElement('a:srgbClr');cc.set('val',LINE if edge=='lnB' else WHITE);sol.append(cc);ln.append(sol);pr.append(ln)
                for p in c.text_frame.paragraphs:
                    p.space_after=Pt(1);p.line_spacing=1.12
                    for r in p.runs:
                        r.font.name=FONT;r.font.size=Pt(size);r.font.color.rgb=color(ACCENT if ri==0 else INK);r.font.bold=ri==0
                        ea=OxmlElement('a:ea');ea.set('typeface',FONT);r._r.get_or_add_rPr().append(ea)
        self.layout.append({'page':len(self.prs.slides),'kind':'table','x':x,'y':y,'w':width,'h':sum(heights),'rows':rows})
        return sh

    def note(self,s,title,seconds,script,ids,extra=''):
        no=len(self.prs.slides);assert all(k in SOURCES for k in ids)
        role='主講' if no<=MAIN_COUNT else '附錄';src='\n'.join(f'[{k}] {SOURCES[k]}' for k in ids)
        text=f'{role} 第{no}頁｜{title}\n\n{script}\n\n補充／證據邊界\n{extra}\n\n來源\n{src}'
        if seconds:text+=f'\n\n預估講解時間：{seconds}秒（非實測配音長度）'
        s.notes_slide.notes_text_frame.text=text
        self.notes.append({'page':no,'kind':role,'title':title,'seconds':seconds,'script':script,'extra':extra,'source_ids':ids,'notes_text':text})

    def build(self):
        E=self.eval['horizons'];D=self.stats['descriptive']
        # 1 Cover.
        s=self.page('','研究提案與初步成果','TW01｜簡潔研究版 v4｜版本日期 2026-10-04')
        self.line(s,.72,1.97,.72,5.57,tint=ACCENT,width=3)
        self.text(s,1.07,2.12,11.35,1.6,'現場菜價7／14天\n預測與分析代理人',36,bold=True)
        self.text(s,1.10,4.19,10.7,1.05,'今天看見的單品價格，\n如何連到一、兩週後的變化？',24,tint=MUTED)
        self.text(s,1.10,5.61,10.7,.60,'AIoT 應用與資料分析\n組員／組長、正式發表日期：待填',14,tint=MUTED)
        self.note(s,'封面',25,'我們的題目是現場菜價七天與十四天預測。使用者在市場看到單品報價，提供今天價格，就能查固定期限的估計；上週同店價可以選填。這次報告包含提案與公開資料的初步結果，同店預測效果仍待驗證。',['P1','STYLE'],'本次改寫只調整論述與呈現，不更換研究題目、模型或資料。')
        # 2 Motivation.
        s=self.page('研究動機：今天的標價，還不足以回答未來','研究背景','[D1–D2] 臺中公開市場歷史資料；比例是歷史配對描述，不是未來機率')
        self.text(s,.75,1.95,5.3,.4,'使用情境',20,bold=True,tint=ACCENT)
        self.text(s,.75,2.67,5.1,.72,'甘藍 32元／台斤',30,bold=True)
        self.text(s,.75,3.49,5.1,1.3,'消費者想知道：\n7天、14天後，會漲、平或跌？',23)
        self.text(s,.75,5.15,5.1,.55,'32元為情境示意；不是本次實際報價。',13,tint=MUTED)
        self.text(s,6.58,1.95,5.8,.86,'歷史樣本中，\n超過±5%變動的配對比例',20,bold=True)
        for i,h in enumerate((7,14)):
            r=D[f'h{h}|ALL'];yy=3.05+i*1.16;bw=3.3*r['percent_exact']/20
            self.text(s,6.58,yy+.04,1.05,.42,f'{h}天',20)
            self.box(s,7.65,yy,bw,.38,ACCENT,ACCENT)
            self.text(s,7.65+bw+.17,yy-.02,1.2,.44,r['percent_display'],22,bold=True,tint=ACCENT)
            self.text(s,7.65,yy+.47,4,.34,f"{r['nonflat']:,}／{r['n']:,}筆歷史配對",13,tint=MUTED)
        self.takeaway(s,'值得測試短期方向，但市場均價不能直接當成同店真值。')
        self.note(s,'研究動機',40,'想像你站在蔬果行，看見甘藍一台斤三十二元。這只是一個今天的價格，還沒有回答一兩週後的變化。右圖用本案歷史資料說明問題確實存在：七天有百分之十點五、十四天有百分之十六點六的配對，變化超過正負百分之五。不過這是歷史市場配對比例，不能當成這家店未來漲跌的機率。因此本案先測公共資料，再把同店效果另列驗證問題。',['D1','D2','P1'],'母體含不同split的市場×品項×起算日×期限配對，並非獨立天數。7天n=10952、14天n=10680；方向按正式標籤含eps容差。')
        # 3 Academic / course prior work.
        s=self.page('相關研究：菜價預測已有方法，場景各不相同','相關研究 1/2','[R01–R02] 論文摘要；[H17] 同課程Team 17原PDF第7–12頁')
        self.table(s,.70,1.88,11.93,[.60,1.06,1.06,1.06],[3.16,5.26,3.51],[
            ['前作','已知做法','值得本案對照'],
            ['陳泯宏，2024','批發歷史價格；比較\nARIMA、ETS、TDNN、LSTM','模型選擇與方向評估'],
            ['李岱鑫，2015','公開資料＋中盤商私人紀錄；\n預測下一週批發價','新增私人資料是否有幫助'],
            ['Team 17，2025\n同課程簡報','市場＋氣象資料；\n各品項LightGBM、未來五天價格','資料、預測期限與應用']],size=19)
        self.takeaway(s,'本案聚焦：現場零售報價、固定7／14天，以及同店到期驗證。')
        self.note(s,'相關研究：預測與私人資訊',40,'我們不是第一個研究菜價的人。陳泯宏比較多種批發價格預測模型；李岱鑫研究加入中盤商私人紀錄是否有價值；同課程Team十七則結合市場和氣象資料，展示未來五天價格。這些都是可比較的前作。本案把問題縮到消費者當下的單店報價、固定七天與十四天，以及如何取得同店到期價格來驗證。期限與輸入不同，不代表本案已經更好。',['R01','R02','H17','P1'],'R01/R02僅依既有摘要查核；R02確切誤差指標仍未知。H17為作者15頁簡報，未重跑其模型；不跨資料比較分數，未載明的功能不等於不存在。')
        # 4 Products: not an absence checklist.
        s=self.page('相關服務：把已知功能與待驗證問題分開','相關研究 2/2','[R05–R06] 已查核官網／商店文案；未安裝實測，文案不是效果證據')
        self.table(s,.70,1.93,11.93,[.63,1.16,1.16,1.24],[2.6,4.75,4.58],[
            ['服務／本案','已見功能或設計','證據界線'],
            ['白菜即時','菜價趨勢預測、零售換算文案','實際模型、固定期限與效果\n依現有證據尚未確認'],
            ['農產分析','輸入購買價、單位換算、\n當日行情比對','其他未載明功能維持未知'],
            ['TW01 本案','今天店價＋選填上週價；\n估計t+7／t+14相對變化','公開市場已有測試；\n同店A／B效果尚未驗證']],size=18,highlight=(3,))
        self.takeaway(s,'差異要靠任務與驗證說清楚，不靠「別人都沒有」來成立。')
        self.note(s,'相關服務與本案缺口',35,'產品端也有前作。白菜即時有趨勢預測與零售換算文案，農產分析已讓人輸入購買價比對當日行情。所以不能說既有工具全部只看今天，或都不能輸入價格。本案要做的是固定期限的估計，並設計同一店家到期價的驗證方式。公開市場測試已經有結果，但這個結果還不能證明某家店的A或B估計比較準。',['R05','R06','P1'],'不對兩項產品未公開或未讀到的其他能力做否定。此頁描述本案提案，不宣稱全球首次。')
        # 5 Questions and contributions.
        s=self.page('研究問題與三項亮點','研究目標','[P1] 本案RQ1–RQ3與貢獻；已完成與待驗證分開')
        self.points(s,[('固定期限的公開評估','比較7／14天方向預測與簡單基線。　已有公開測試'),('今天價與上週價的A／B設計','沒有上週價仍可用；新增資訊是否有幫助。　待同店驗證'),('看得懂、查得到來源的結果','同時呈現期限、價格估計與資料範圍。　原型已展示')],y=1.92,gap=1.35,size=23)
        self.note(s,'研究問題與三項亮點',35,'這三項亮點各自對應一個問題。第一，固定七天與十四天時，模型相較不變價和動量表現如何。第二，使用者多提供一筆可比較的上週同店價格，是否有實際幫助。第三，結果能不能讓人看懂期限、來源與限制。第一項已有公開資料的測試；第二項等真店到期驗證；第三項有原型，但使用者理解與輸入負擔還沒有實測。',['P1','E1'],'不把C1已有結果當成原創性或全面泛化已證實，也不把C3原型當成完整保存重播驗收。')
        # 6 Architecture: independent stores and one short subtitle per module.
        s=self.page('系統架構：輸入、計算與資料各有角色','提出的方法','[A1] 依本機原型整理；編號是功能索引，不是固定執行順序')
        for x,y,w,h,head,body in [(.75,2.05,3.2,1.35,'① 使用者與手機','品項、報價、日期'),(5.05,2.05,3.2,1.35,'② 查詢處理','確認條件與單位'),(9.25,2.05,3.25,1.35,'③ 菜價計算','模型推論＋A／B換算'),(9.25,4.70,3.25,1.15,'④ 公開歷史','市場均價・唯讀'),(5.05,4.70,3.2,1.15,'⑤ 預測紀錄','自願保存・獨立儲存')]:
            self.box(s,x,y,w,h);self.text(s,x+.15,y+.25,w-.30,.43,head,21,bold=True,align='center');self.text(s,x+.15,y+.80,w-.30,.40,body,17,tint=MUTED,align='center')
        self.line(s,3.95,2.48,5.05,2.48,tint=ACCENT,width=1.6,arrow=True);self.text(s,4,2.03,1,.34,'提交',13,align='center',tint=ACCENT)
        self.line(s,5.05,3.15,3.95,3.15,tint=ACCENT,width=1.6,arrow=True);self.text(s,3.98,3.48,1.08,.4,'回傳結果',13,align='center',tint=ACCENT)
        self.line(s,8.25,2.48,9.25,2.48,tint=ACCENT,width=1.6,arrow=True);self.text(s,8.27,2.03,.96,.34,'計算',13,align='center',tint=ACCENT)
        self.line(s,9.25,3.15,8.25,3.15,tint=ACCENT,width=1.6,arrow=True);self.text(s,8.25,3.48,1,.4,'回傳預測',13,align='center',tint=ACCENT)
        self.line(s,10.86,3.4,10.86,4.7,tint=ACCENT,width=1.6,arrow=True);self.text(s,11.03,3.89,1.2,.43,'讀取',14,tint=ACCENT)
        self.line(s,5.98,3.4,5.98,4.7,tint=ACCENT,width=1.6,arrow=True);self.text(s,4.88,3.9,1,.65,'選擇時\n保存',13,tint=ACCENT,align='right')
        self.line(s,7.30,4.7,7.30,3.4,tint=ACCENT,width=1.6,arrow=True);self.text(s,7.45,3.9,.85,.43,'查閱',13,tint=ACCENT)
        self.text(s,.75,4.85,3.65,.85,'公開歷史與私人紀錄分開，\n不把個人報價寫進公開資料。',17,tint=MUTED)
        self.takeaway(s,'本頁呈現目前原型；語言模型整合與外網部署屬後續工作。')
        self.note(s,'系統架構',45,'請按編號看這張圖。第一，使用者提供品項和現場報價。第二，查詢層確認日期、單位及上週資訊是否可比較，再呼叫第三個菜價計算工具。計算工具讀取第四個公開歷史資料，產生七天和十四天估計，結果經查詢層回到手機。只有使用者選擇保存，才把當時結果寫入第五個預測紀錄。公開資料與私人紀錄是分開的。現在講的是本機原型，語言模型後端和外網服務仍是後續整合。',['A1','M2','P1'],'①→②提交、②→③計算、③→④讀取、③→②→①回傳；②与⑤保存/查閱，④⑤無直接寫入。GLM/Qwen是語言模型後端；Hermes是暫定流程框架；ForecastEngine才計算菜價。本次不驗收或修復既有agent缺陷。')
        # 7 Data.
        s=self.page('資料來源與整理','資料與方法','[D1] 臺中市政府公開資料快照；[D2] 既有資料品質與標籤統計')
        self.table(s,.75,1.9,11.85,[.61,.71,.73,.71,.71,.74],[2.62,9.23],[['項目','本案使用範圍'],['來源','臺中市公有零售市場每日訪價均價'],['期間／市場','2025-10-04～2026-10-02；排除停用市場後13市場'],['品項','甘藍、青江白菜、絲瓜'],['單位／缺值','元／台斤；0代表無報價，不是零元'],['整理方式','統一日期與單位；建立同市場、同品項的到期配對']],size=21)
        self.takeaway(s,'一筆資料代表市場訪價均價，不代表其中某一家店的成交價。')
        self.note(s,'資料來源與整理',35,'本案使用臺中市公有零售市場的每日訪價均價，期間約一年，排除停用市場後共十三個市場。第一版先做甘藍、青江白菜與絲瓜。單位是元每台斤，欄位中的零表示沒有報價，所以不能直接當成零元拿去訓練。整理後依同市場、同品項、同日期配對，建立未來七天和十四天的目標。這裡沒有單店識別，因此公開測試和單店效果必須分開解釋。',['D1','D2'],'原始4450列×63欄；原三品項長表13350列，0值905格轉缺值。13市場中甘藍有效配對涵蓋12市場。364是原檔不同日期數，不表示每市場每日無缺。')
        # 8 Target.
        s=self.page('預測目標：從報價日算起的7／14天變化','資料與方法','[M1] 固定日曆期限；±5%含邊界為持平；下方數值均為教學例')
        for x,head,sub in [(1,'t','本次報價日'),(6,'t＋7','7個日曆日後'),(10.35,'t＋14','14個日曆日後')]:
            self.text(s,x,1.90,2,.5,head,26,bold=True,tint=ACCENT,align='center');self.text(s,x-.35,2.52,2.7,.42,sub,18,align='center')
        self.line(s,2,3.2,11.35,3.2,tint=ACCENT,width=1.8,arrow=True)
        for x in [2,7,11.35]:self.line(s,x,3.07,x,3.34,tint=ACCENT,width=1.8)
        self.text(s,1,3.76,11.35,.63,'相對變化 ＝（到期價格 − 當日價格）÷ 當日價格',24,bold=True,align='center')
        self.table(s,.90,4.7,11.50,[.58,.77],[3.83,3.84,3.83],[['下跌','持平','上漲'],['小於 −5%','介於 −5% 與 ＋5%','大於 ＋5%']],size=21)
        self.text(s,.95,6.34,11.4,.48,'例如市場均價40→43元：＋7.5%，上漲；40→42元：＋5%，持平。',19,tint=ACCENT)
        self.note(s,'預測目標與漲跌定義',40,'我們不預測模糊的下週平均，而是從本次可信報價日算起，精確七天與十四個日曆日。相對變化就是到期價格減去當日價格，再除以當日價格。超過正負百分之五才算上漲或下跌，剛好落在百分之五仍屬持平。像四十元變四十三元，是上漲百分之七點五；變四十二元剛好百分之五，仍屬持平。公開資料用市場均價建標籤，同店驗證則要用同店到期價。',['M1'],'正式分類用eps=1e-9處理浮點邊界。t不是公開行情最後日期，也不改成交易日。公開標籤m(t+h)/m(t)-1；同店真值另收，不以市場值代填。')
        # 9 A.
        s=self.page('A模式：沒有上週價，也能查','核心方法 1/2','[M2] 現價錨定方法；本頁全部數值為教學假設，非真店實驗或模型fixture')
        self.text(s,.8,1.89,11.7,.52,'今天店價 ×（1＋模型估計的相對變化）＝ 到期價格估計',23,bold=True)
        self.text(s,.8,2.78,11.7,.53,'假設今天32元／台斤；模型的7天估計為 −4.1%。',22)
        self.text(s,.8,3.5,11.7,.64,'32 ×（1 − 0.041）＝ 30.688 ≈ 30.7元',27,bold=True,tint=ACCENT)
        self.table(s,.82,4.62,11.69,[.55,.60,.60],[3.2,3.2,3,2.29],[['期限','假設相對變化','價格估計','方向'],['7天','−4.1%','30.7元','持平'],['14天','−0.9%','31.7元','持平']],size=19)
        self.text(s,.83,6.61,11.7,.32,'輸入現價改變金額尺度，不會讓方向預測自動變準。',19,bold=True,tint=ACCENT)
        self.note(s,'A模式',45,'A模式不需要上週的紀錄。假設今天店價是三十二元，模型估計七天相對變化為負百分之四點一，我們把三十二乘以零點九五九，得到三十點六八八元，顯示約三十點七元。雖然金額下降，但還在正負百分之五內，所以分類是持平。下方十四天也是同樣算法。這些都是教學假設，不是本次查詢的真模型輸出。現價錨定只改顯示金額，並不能證明方向變得更準。',['M2','M1'],'A公式p_hat=p*(1+r_hat)。p>0時，[p(1+r_hat)-p]/p=r_hat；同一分類器方向不變。14天32*0.991=31.712≈31.7；負號不等同down。')
        # 10 B.
        s=self.page('B模式：多一筆上週價，是否真的有幫助？','核心方法 2/2','[M2] α＝0.3為固定實驗規則；本頁所有價格與變化數字均為教學假設')
        self.text(s,.8,1.87,11.72,.67,'在A的基礎上，加上「店價變化與公開市場變化的差」。',23,bold=True)
        self.table(s,.8,2.85,6.1,[.60,.64,.64,.64],[3,3.1],[['教學條件','數值'],['今天／上週店價','32元／28元'],['店價週變化','32 ÷ 28 − 1 ≈ 14.29%'],['公開同期變化','假設為 ＋1%']],size=18)
        self.text(s,7.48,2.84,4.87,.5,'B的7天調整',22,bold=True,tint=ACCENT)
        self.text(s,7.48,3.66,4.8,1.58,'−4.1%\n＋0.3 ×（14.29% − 1%）\n≈ −0.11%　仍為持平',21)
        self.text(s,7.48,5.37,4.7,.65,'同店、同品項、同規格；\n日期與單位需可比較。',17,tint=MUTED)
        self.takeaway(s,'係數0.3不是學得的效果；B是否更好，要等同店到期價驗證。')
        self.note(s,'B模式',45,'有可比較的上週同店價時，才使用B。這個例子假設今天三十二元、上週二十八元，店價一週變化約百分之十四點二九；公開同期假設是百分之一。B把兩者的差乘以零點三，加回A的負百分之四點一，得到約負百分之零點一一，仍屬持平。請注意，零點三只是固定實驗規則，不是從同店資料學來的效果。B能不能更好，要把同一批有上週價的個案，同時保留A和B，再等到期真值比較。',['M2','P1'],'r_B=r_A+0.3*(Delta_store-Delta_pub)；本例未四捨五入r_B≈-0.001142857。上週可比日期7±2天，同店/品項/規格，單位需轉同尺度；缺少或不可比則僅A，不用0或公開t-7冒充私人資料。')
        # 11 Model selection.
        s=self.page('先用驗證集選模型，再看封存測試','評估設計','[M3, S1] 保留樣本涵蓋期；成熟清除規則及四段範圍見附錄A')
        for x,w,label,span in [(.84,3.76,'訓練','2025-10-04～2026-05-08'),(4.76,3.76,'驗證','2026-05-10～2026-07-21'),(8.68,3.76,'測試','2026-07-22～2026-10-02')]:
            self.text(s,x,1.95,w,.49,label,23,bold=True,tint=ACCENT,align='center');self.line(s,x,2.65,x+w,2.65,tint=ACCENT,width=6);self.text(s,x-.05,3,w+.1,.45,span,16,align='center')
        self.table(s,.85,4.03,11.60,[.57,.72,.72],[2.45,5,4.15],[['期限','選定模型','選擇方式'],['7天','隨機森林（RF，300棵樹）','依驗證集macro-F1選型'],['14天','梯度提升（HGB，200輪）','再以同一封存測試列比較基線']],size=20)
        self.takeaway(s,'特徵只使用當下可見的歷史；邊界樣本依到期可見性清除。')
        self.note(s,'模型選擇與評估設計',40,'評估的先後順序很重要。我們先用較早資料訓練，再用驗證集的macro-F1選模型，最後才看保留的測試資料。七天選到隨機森林，十四天選到梯度提升。模型輸入是十二個數值特徵，再加品項編碼，不是另外訓練一個消費者語言模型。圖上是保留樣本的日期範圍，邊界上如果到期價格在下一段才知道，就要清除相關樣本。資料延遲目前保守假設一天，細節放在附錄。',['M3','S1'],'同日起算觀測不隨機分到不同時間區；依h分別清除。時間圖不是所有日期所有期限皆保留的宣稱。測試n(h7)=2085、n(h14)=1859；其他有效標籤未必符合特徵和測試範圍。')
        # 12 Direction scores.
        s=self.page('結果：方向分數仍未超過動量基線','初步結果 1/2','[E1] 封存測試：7天 n＝2,085；14天 n＝1,859；不同方法使用相同測試列')
        self.text(s,.82,1.91,11.7,.55,'方向macro-F1：上漲、持平、下跌分別計分後平均，越高越好。',20)
        self.table(s,.82,2.94,11.69,[.64,.82,.82,.88],[6.19,2.75,2.75],[['方法','7天','14天'],['不變價：全部預測持平',f"{E['h=7']['baselines']['flat_no_change']['direction']['macro_f1']:.4f}",f"{E['h=14']['baselines']['flat_no_change']['direction']['macro_f1']:.4f}"],['動量：過去同長度變化延續',f"{E['h=7']['baselines']['past_change_continues']['direction']['macro_f1']:.4f}",f"{E['h=14']['baselines']['past_change_continues']['direction']['macro_f1']:.4f}"],['本案選定模型',f"{E['h=7']['selected']['direction']['macro_f1']:.4f}",f"{E['h=14']['selected']['direction']['macro_f1']:.4f}"]],size=23,highlight=(3,))
        self.takeaway(s,'本案高於全持平，但兩個期限都未勝動量；不能只挑較弱基線。')
        self.note(s,'方向結果',45,'這張表用同一批封存測試列比較三種方法。不變價永遠預測持平；動量延續過去同長度的變化；最後一列是本案選定模型。七天模型約零點四二一八，十四天約零點四一六九，都高於全部持平，但都沒有超過動量。這表示不能只報告比全持平好，就說模型已經成功。這裡也沒有做顯著性檢定，不能把小數差異說成統計上顯著勝出。',['E1'],'macro-F1是各類別F1非加權平均，不是準確率百分比，也不等同個店效果。RF/HGB不是LLM猜價。表格直接讀封存JSON。')
        # 13 MAE and missed rises.
        s=self.page('結果限制：價格誤差與上漲事件仍需改善','初步結果 2/2','[E1] 相對變化MAE越小越好；數值不是元，也不是準確率')
        self.table(s,.85,1.93,11.6,[.64,.84,.84],[6.1,2.75,2.75],[['相對變化MAE','7天','14天'],['不變價',f"{E['h=7']['baselines']['flat_no_change']['return_mae']:.4f}",f"{E['h=14']['baselines']['flat_no_change']['return_mae']:.4f}"],['本案選定模型',f"{E['h=7']['selected']['return_mae']:.4f}",f"{E['h=14']['selected']['return_mae']:.4f}"]],size=23)
        self.text(s,.86,4.78,11.6,.61,'7天實際上漲158筆，模型抓到16筆（召回10.1%）。',23,bold=True)
        self.text(s,.86,5.56,11.6,.57,'方向與金額各有不足；目前不宜把它當成可靠的上漲警報。',21,tint=MUTED)
        self.takeaway(s,'先呈現能力邊界，再用真店與後續實驗判斷是否值得使用。')
        self.note(s,'誤差與上漲召回',40,'除了方向，我們也比較相對變化的平均絕對誤差。這個指標越小越好，但七天與十四天的模型誤差都比不變價高。再看上漲事件，七天實際有一百五十八筆上漲，模型只抓到十六筆，召回約百分之十點一。所以目前還不能把結果當成可靠的上漲警報，更不能保證依預測採買就省錢。這些負向結果會限制產品用途，也決定後續要驗證什麼。',['E1'],'h14上漲召回23/268≈8.6%。MAE為相對變化誤差，0.0664約6.64個百分點的變化量誤差，不等於價格百分比誤差或命中率。')
        # 14 Screenshot: exact historical pixels, enlarged 7-day area.
        s=self.page('介面示例：同一期限，對照A與B','原型展示','[U1] 2026-10-03本機原型，放大上方7天結果區；非本次即時查詢')
        cropped=self.evidence/'ui_7day_crop.png';im=Image.open(self.shot);im.crop((0,0,390,590)).save(cropped)
        s.shapes.add_picture(str(cropped),Inches(.95),Inches(1.77),height=Inches(4.95))
        self.points(s,[('先看期限','確認這組結果是7天，另一組才是14天。'),('再看A／B','對照不使用與使用上週資訊的結果。'),('最後讀範圍','市場均價、店價估計與實驗性提醒分開。')],x=5.35,y=2.1,w=7,size=22,gap=1.38)
        self.note(s,'介面示例',40,'這是既有本機原型的歷史截圖，不是今天重新查詢的結果。為了讓字能看清楚，這頁放大上方七天的A與B。讀畫面時先看期限，再看使用或不使用上週資訊的差異，最後看來源與實驗性提醒。畫面中的B與前面教學例數值不同，因為這張圖來自另一個歷史查詢，不能混成同一筆個案。現在可以展示介面，但完整保存重播、語言模型後端及外網服務仍未完成驗收。',['U1','A1','P1'],'圖片只裁切放大7天比較；原圖390×844保留原位，不改數值或日期。截圖不證明自然語言後端、保存或即時服務已驗收。')
        # 15 Plan.
        s=self.page('下一步：用對應證據回答三個問題','驗證計畫','[P1, C1] 設計藍圖與規格管理：design.md＋OpenSpec；實際日期待定')
        self.table(s,.8,1.94,11.75,[.66,1.08,1.08,1.08],[2.86,6.60,2.29],[['要驗證什麼','需要的證據','目前'],['同店A／B增益','先保存A與B；7／14天後取得同店同規格價格','待執行'],['低負擔與理解','輸入時間、錯誤率；能否說明期限與資料範圍','待執行'],['服務是否可用','後端整合、外網手機查詢及端到端驗證','暫緩']],size=20)
        self.takeaway(s,'公開測試、同店驗證與系統可用性，是三種不同的證據。')
        self.note(s,'後續驗證',35,'後續驗證不只再訓練一次。第一，保存同一筆查詢的A與B，等到七天、十四天取得同店同規格真價，才比較輸入資訊的增益。第二，觀察使用者是否能順利輸入，並正確理解資料範圍。第三，後端與外網服務另做端到端驗證。這三類證據不能互相代替。設計藍圖與OpenSpec管理保留在專案中，目前日期還沒排定，這次也不啟動暫緩的整合工作。',['P1','C1','M2'],'同筆A/B配對，先凍結再取到期真值；不用公開t-7或均價代替私人到期資料。影片和發布獨立追蹤，本次不發布。')
        # 16 Conclusion.
        s=self.page('結論與後續工作','結論','TW01｜研究提案與初步成果｜主講結束；後附計算、數據與引用')
        self.points(s,[('已完成','公開市場7／14天評估，以及現場報價的A／B原型。'),('目前限制','方向未勝動量、誤差未勝不變價；同店效果尚未驗證。'),('下一步','取得同店到期真值，判斷上週資訊能否帶來實際增益。')],y=2.03,gap=1.39,size=24)
        self.note(s,'結論',30,'總結來說，我們把現場報價連到固定七天與十四天的預測任務，完成公開市場評估與A/B原型。已有結果也清楚顯示限制：方向未超過動量，金額誤差未超過不變價，同店效果還沒有證據。下一步不是先保證更準，而是取得同店到期真值，回答多一筆上週價格到底有沒有幫助。謝謝，後面的附錄備有計算與資料供提問。',['P1','E1'],'本次為作者重寫與自查，非另一次獨立研究驗收；不更換題目或模型。')
        # 17 Appendix A.
        s=self.page('附錄A｜資料口徑與成熟清除','附錄','[D2, S1] 既有標籤與切分資料；本版唯讀使用')
        rows=[['邊界／期限','清除起算日範圍','列數']]
        for boundary,label in [('train_to_val','訓練→驗證'),('val_to_test','驗證→測試')]:
            for h in ['h7','h14']:
                r=self.stats['purge']['segments'][boundary][h];rows.append([f'{label}／{h[1:]}天',f"{r['t_min']}～{r['t_max']}",str(r['rows'])])
        self.table(s,.8,1.86,11.73,[.6,.61,.61,.61,.61],[3.55,5.86,2.32],rows,size=18)
        self.text(s,.86,5.22,11.65,1.05,'成熟日＝t＋h＋1；四段合計1,386列，聯集30個起算日。\n7天有效標籤10,952筆；14天10,680筆，與測試樣本數不同。',19)
        self.note(s,'附錄A 資料與成熟清除',0,'此表回答成熟清除的日期與母體。7天對應t+8、14天對應t+15，分別清除跨越下一區界的標籤。不把總體起算日期範圍當成所有期限共用同一清除帶。',['D1','D2','S1'],'資料品質：原3品項0值905格，整份原始不同日期364；全61價格欄零值比例與前三品項905不同母體。21筆邊界來自直接abs(r)>0.05與正式eps方向口徑差異，不改labels。')
        # 18 Appendix B.
        s=self.page('附錄B｜A／B公式與可比條件','附錄','[M1–M2] 使用固定α＝0.3，不宣稱係數已學得或已驗證有增益')
        self.text(s,.84,1.93,11.6,.53,'A：p̂(t＋h) ＝ p(t) ×（1＋r̂_h）',24,bold=True)
        self.text(s,.84,2.81,11.6,.59,'B：r_B ＝ r_A ＋ 0.3 ×（Δ店價 − Δ公開同期）',24,bold=True)
        self.text(s,.87,3.85,11.6,1.26,'p(t)：今天店價；r_A：公開模型的相對變化估計。\nΔ店價：今天價 ÷ 上週價 − 1；公開同期使用可比期間。\n上週價要有實際日期，並確認同店、同品項、同規格與單位。',19)
        self.text(s,.87,5.6,11.6,.96,'p＞0時，[p(1＋r̂)−p]／p＝r̂。\n同一分類器下，A錨定只改金額；缺少可比上週資訊則回A。',20,tint=ACCENT)
        self.note(s,'附錄B 公式',0,'A的方向不變是代數性質，不是準確率實驗。B加上差動量後可以改變相對變化，但是否更接近真店到期價仍要驗證。',['M1','M2'],'7±2天是目前可比上週資訊的日期契約；單位先轉同尺度。方法設計不代表每一產品入口校驗均已通過驗收。')
        # 19 Appendix C.
        s=self.page('附錄C｜分品項方向結果','附錄','[E2] 公開市場封存測試；不以事件比例解釋分數差異的因果')
        rows=[['品項','7天：n／macro-F1','14天：n／macro-F1']]
        for crop,cn in [('cabbage','甘藍'),('bok_choy','青江白菜'),('sponge_gourd','絲瓜')]:
            a=self.eval['by_crop'][f'h=7|{crop}'];b=self.eval['by_crop'][f'h=14|{crop}'];rows.append([cn,f"{a['n']:,}／{a['macro_f1']:.4f}",f"{b['n']:,}／{b['macro_f1']:.4f}"])
        self.table(s,.85,2.12,11.6,[.7,.89,.89,.89],[3,4.3,4.3],rows,size=24)
        self.text(s,.87,6,11.6,.61,'分數隨品項不同；此處只描述結果，不推定差異原因。',21,tint=ACCENT)
        self.note(s,'附錄C 分品項',0,'全部使用封存JSON by_crop，不跨資料排名。市場平均價表現不能當成任何一店的準確度。',['E2'],'同品項時序與市場觀測可能相關，未做顯著性或成因檢定。')
        # 20 Appendix D.
        s=self.page('附錄D｜7天混淆矩陣：上漲容易漏掉','附錄','[E1] 列為實際類別，欄為模型預測類別；n＝2,085')
        d=E['h=7']['selected']['direction'];names=['下跌','持平','上漲'];rows=[['實際＼預測',*names]]
        rows += [[names[i],*map(str,row)] for i,row in enumerate(d['confusion_matrix'])]
        self.table(s,.85,2.14,7.45,[.76,.88,.88,.88],[2.2,1.75,1.75,1.75],rows,size=24)
        self.text(s,9,2.24,3.5,.5,'上漲召回',23,bold=True,tint=ACCENT)
        self.text(s,9,3.02,3.5,1.08,'16 ÷ 158\n≈ 10.1%',27,bold=True)
        self.text(s,9,4.62,3.45,1.1,'142筆實際上漲，\n沒有被報為上漲。',20)
        self.note(s,'附錄D 混淆矩陣',0,'實際上漲的一百五十八筆裡，十六筆預測上漲；其餘四十三筆預測下跌，九十九筆預測持平，合計漏報一百四十二筆。',['E1'],'這是7天RF的既有封存測試，不是新訓練結果；類別順序依JSON down/flat/up。')
        # 21 Appendix E.
        s=self.page('附錄E｜主要引用與同課程前作','附錄','完整網址、讀取層級與本機定位見本頁備註及同版SOURCES.md')
        entries=[('[R01] 陳泯宏，2024','國產大宗蔬菜價格預測模型之比較。臺大碩論；已讀摘要。'),('[R02] 李岱鑫，2015','高雄果菜運銷公司高麗菜批發價格預測之研究—蒐集私有資料的價值。已讀摘要。'),('[H17] 同課程Team 17，2025','蔬果價格預測：LightGBM的最佳採購指南。15頁PDF；僅作者簡報。'),('[R05] 白菜即時；[R06] 農產分析','官網／商店功能文案；未安裝實測。'),('[D1] 臺中市政府公開資料','公有零售市場每日蔬果價格表。data.gov.tw／84539。')]
        for i,(head,body) in enumerate(entries):
            y=1.83+i*.95;self.text(s,.82,y,11.68,.35,head,18,bold=True,tint=ACCENT);self.text(s,.82,y+.39,11.68,.51,body,16)
        self.note(s,'附錄E 引用',0,'前作只引用本案已讀證據支持的做法；摘要未載明功能不視為不存在。Team17為同課程作者簡報，不是本案重現的性能證據。',['R01','R02','H17','R05','R06','D1'],'只借用歷屆資訊安排與論述方法，不複製他組圖片或實驗成果。所有本案分數依封存結果。')
        # 22 Appendix F: implementation details belong here.
        s=self.page('附錄F｜實作與交付索引','附錄','[A1, P1, C1] 技術定位供備詢；不把工具安裝或文件存在等同全案驗收')
        self.table(s,.82,1.93,11.69,[.62,.79,.79,.79,.89],[3.1,5.6,2.99],[['部分','定位','狀態'],['查詢與計算','/forecast、/agent/query；ForecastEngine','本機原型'],['選擇性保存','/records；private/predictions','完整重播未驗收'],['資料與模型','public_history；RF／HGB模型包','既有封存工件'],['研究管理','design.md＋OpenSpec；來源與評估紀錄','文件與規格已建']],size=18)
        self.text(s,.86,6.08,11.55,.71,'GLM／Qwen是語言模型後端；Hermes是流程框架。\n語言模型整合、外網部署與正式發布：仍屬後續工作。',18,tint=MUTED)
        self.note(s,'附錄F 實作與交付',0,'完整端點、檔名、來源hash及測試記錄保留在專案，主講只留理解研究所需資訊。本次重寫與自查不替代產品代理人的獨立驗收。',['A1','P1','C1','STYLE'],'不修改design.md等既有研究文件，也不宣稱前輪跨文件問題全面結案。新版PPT與講稿使用一致角色。影片只有兩頁試片，本次不擴製或上傳。')
        assert len(self.prs.slides)==22 and len(self.notes)==22

    def save(self):
        self.out.mkdir(parents=True,exist_ok=True);ppt=self.out/(NAME+'.pptx')
        if ppt.exists():raise FileExistsError(ppt)
        self.prs.save(str(ppt));temp=ppt.with_suffix('.tmp')
        with zipfile.ZipFile(ppt) as zin,zipfile.ZipFile(temp,'w',zipfile.ZIP_DEFLATED) as zout:
            for n in sorted(zin.namelist()):
                zi=zipfile.ZipInfo(n,date_time=(2026,10,4,0,0,0));zi.compress_type=zipfile.ZIP_DEFLATED;zout.writestr(zi,zin.read(n))
        temp.replace(ppt)
        total=sum(r['seconds'] for r in self.notes)
        records={'version':VERSION,'title':self.prs.core_properties.title,'main_slides':16,'appendix_slides':6,'estimated_main_seconds':total,'time_note':'講解設計估時，不是實際配音或錄影長度。','slides':self.notes}
        (self.out/'notes_v4.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
        lines=['# TW01 v4 逐頁講稿','',f'配套：{NAME}.pptx；16頁主講＋6頁附錄。',f'主講設計估時：{total//60}分{total%60:02d}秒；尚未配音或錄製。','','PPT備註與本檔同源；详细引用見每頁備註與SOURCES.md。','']
        for r in self.notes:
            lines.extend([f"## {r['kind']} 第{r['page']}頁｜{r['title']}",'',r['script'],'','**補充／證據邊界**：'+r['extra'],'','**來源**：'+'、'.join('['+x+']' for x in r['source_ids']),''])
        (self.out/'SLIDE_NOTES_v4.md').write_text('\n'.join(lines),encoding='utf-8')
        (self.out/'SOURCES.md').write_text('# TW01 v4 來源定位\n\n'+'\n\n'.join(f'## [{k}]\n{v}' for k,v in SOURCES.items()),encoding='utf-8')
        (self.evidence/'layout_manifest.json').write_text(json.dumps(self.layout,ensure_ascii=False,indent=2),encoding='utf-8')
        inputs=['outputs/evaluation/public/public_test_results.json','outputs/proposal_delivery/run3_20261004_01/stats_master.json','outputs/e2e/ui_ab_mobile_390x844.png']
        (self.evidence/'build_manifest.json').write_text(json.dumps({'version':VERSION,'pptx':str(ppt),'slides':22,'main':16,'appendix':6,'seconds_estimated':total,'source_sha256':{p:sha(self.root/p) for p in inputs},'pptx_sha256':sha(ppt),'style':'white, short titles, one primary question per slide, native editable shapes/tables','limits':'No model inference/training, product tests, animation, video, publishing or independent acceptance.'},ensure_ascii=False,indent=2),encoding='utf-8')
        return ppt

def main():
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('--root',type=Path,required=True);ap.add_argument('--out',type=Path,required=True);ap.add_argument('--evidence',type=Path,required=True)
    a=ap.parse_args();a.evidence.mkdir(parents=True,exist_ok=True)
    d=Deck(a.root.resolve(),a.out.resolve(),a.evidence.resolve());d.build();p=d.save()
    print(json.dumps({'pptx':str(p),'pages':len(d.prs.slides),'notes':len(d.notes),'estimated_seconds':sum(r['seconds'] for r in d.notes)},ensure_ascii=False))

if __name__=='__main__':main()
