# -*- coding: utf-8 -*-
"""Check the authored v4 deck against its frozen sources and same-version notes."""
from __future__ import annotations
import argparse
import hashlib
import json
from pathlib import Path
from pptx import Presentation
from pptx.oxml.ns import qn
from pypdf import PdfReader

def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def norm(text):return ''.join(text.split())
def texts(slide):
    values=[]
    for sh in slide.shapes:
        if sh.has_text_frame:values.append(sh.text)
        if sh.has_table:values.extend(c.text for row in sh.table.rows for c in row.cells)
    return values

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--root',required=True,type=Path);ap.add_argument('--package',required=True,type=Path);ap.add_argument('--render',required=True,type=Path);ap.add_argument('--out',required=True,type=Path)
    a=ap.parse_args();root=a.root.resolve();out=a.package.resolve();checks=[]
    def check(name,condition,detail=None):checks.append({'check':name,'passed':bool(condition),'detail':detail})
    pptx=next(out.glob('*.pptx'));pdf=pptx.with_suffix('.pdf');p=Presentation(str(pptx));q=PdfReader(str(pdf));notes=json.loads((out/'notes_v4.json').read_text(encoding='utf-8'));md=(out/'SLIDE_NOTES_v4.md').read_text(encoding='utf-8');refs=(out/'SOURCES.md').read_text(encoding='utf-8');ev=json.loads((root/'outputs/evaluation/public/public_test_results.json').read_text(encoding='utf-8'));st=json.loads((root/'outputs/proposal_delivery/run3_20261004_01/stats_master.json').read_text(encoding='utf-8'))
    check('ppt_pdf_notes_are_22',len(p.slides)==len(q.pages)==len(notes['slides'])==22)
    check('main_16_appendix_6',sum(x['kind']=='主講' for x in notes['slides'])==16 and sum(x['kind']=='附錄' for x in notes['slides'])==6)
    check('source_versions_v4',notes['version']=='v4' and all('v2' not in x and 'v3' not in x for s in p.slides for x in texts(s)))
    for s,r in zip(p.slides,notes['slides']):
        check('notes_page_'+str(r['page']),s.has_notes_slide and norm(s.notes_slide.notes_text_frame.text)==norm(r['notes_text']) and norm(r['script']) in norm(md) and all('['+key+']' in refs for key in r['source_ids']))
    for h in ['h=7','h=14']:
        e=ev['horizons'][h]
        actual=' '.join(texts(p.slides[11]))
        expected=[e['selected']['direction']['macro_f1'],e['baselines']['flat_no_change']['direction']['macro_f1'],e['baselines']['past_change_continues']['direction']['macro_f1']]
        check('direction_scores_'+h,all(f'{n:.4f}' in actual for n in expected))
        actual=' '.join(texts(p.slides[12]));expected=[e['selected']['return_mae'],e['baselines']['flat_no_change']['return_mae']]
        check('mae_scores_'+h,all(f'{n:.4f}' in actual for n in expected))
    motivation=' '.join(texts(p.slides[1]))
    check('motivation_counts',all(str(v) in motivation for v in ['10.5%','16.6%','1,149','10,952','1,773','10,680']))
    teaching=' '.join(texts(p.slides[8])+texts(p.slides[9]))
    check('teaching_values_labelled','全部數值為教學假設' in teaching and '30.688' in teaching and '−0.11%' in teaching)
    check('frozen_eval_unchanged',sha(root/'outputs/evaluation/public/public_test_results.json')=='2d28acd2f2c3e5ce566b9095c6a85477db8faf858d9ce29e974caaedc25674b8')
    purge=' '.join(texts(p.slides[16]));check('purge_dates_and_rows',all(str(v) in purge for b in st['purge']['segments'].values() for r in b.values() for v in [r['t_min'],r['t_max'],r['rows']]))
    diagram=' '.join(texts(p.slides[5]));check('numbered_function_groups',all(x in diagram for x in ['① 使用者與手機','② 查詢處理','③ 菜價計算','④ 公開歷史','⑤ 預測紀錄']))
    check('diagram_relations',len(p.slides[5].shapes._spTree.findall(qn('p:cxnSp')))>=7 and all(x in diagram for x in ['讀取','查閱','保存','回傳']))
    pictures=[(i+1,sh) for i,s in enumerate(p.slides) for sh in s.shapes if sh.shape_type==13]
    tables=sum(sh.has_table for s in p.slides for sh in s.shapes)
    check('editable_not_flattened',len(pictures)==1 and pictures[0][0]==14 and tables>=10,{'pictures':len(pictures),'tables':tables})
    outside=[]
    for i,s in enumerate(p.slides,1):
        for sh in s.shapes:
            if sh.left<0 or sh.top<0 or sh.left+sh.width>p.slide_width+100 or sh.top+sh.height>p.slide_height+100:outside.append({'page':i,'name':sh.name})
    check('objects_inside_slide',not outside,outside)
    effects=[]
    for i,s in enumerate(p.slides,1):
        for sh in s.shapes:
            if sh._element.findall('.//'+qn('a:outerShdw')):effects.append(i)
            for ref in sh._element.findall('.//'+qn('a:effectRef')):
                if ref.get('idx')!='0':effects.append(i)
    check('no_theme_shadow',not effects,effects)
    check('no_animations',all(s._element.find(qn('p:timing')) is None for s in p.slides))
    meta=json.loads((a.render/'render_meta.json').read_text(encoding='utf-8'));check('actual_office_export',meta['page_count_match'] and meta['source_unchanged'] and not meta['application_quit_called'])
    check('final_plain_text_height',not [x for x in meta['text_height_candidates'] if 'overflow_pt' in x])
    check('pdf_has_extractable_text',all(len(pg.extract_text().strip())>10 for pg in q.pages))
    result={'passed':all(x['passed'] for x in checks),'checks':checks,'pptx_sha256':sha(pptx),'pdf_sha256':sha(pdf),'limits':'Author checks, not independent review. Table text geometry in Office was unavailable; table readability checked visually. No new model experiments or whole-project acceptance.'}
    a.out.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'passed':result['passed'],'checks':len(checks),'failures':[x for x in checks if not x['passed']]},ensure_ascii=False,indent=2))
    if not result['passed']:raise SystemExit(1)

if __name__=='__main__':main()
