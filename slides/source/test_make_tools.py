# -*- coding: utf-8 -*-
"""D44-R06 製作工具測試（放製作目錄，不動產品 tests/）。

執行：.venv/Scripts/python.exe -m pytest slides/source/test_make_tools.py -q
覆蓋：來源表數值、slide_id 映射、MD/PPT notes 同源（正規化包含）、缺圖明確失敗、防覆寫、頁數一致。
"""
import json
import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
RUN2 = ROOT / "outputs/proposal_delivery/run2_20261004_01"
PPTX_V2 = ROOT / "slides/TW01_菜價預測_研究提案與初步成果_v2.pptx"
PDF_V2 = PPTX_V2.with_suffix(".pdf")
NOTES_JSON = ROOT / "slides/source/notes_v2.json"
STATS = RUN2 / "stats_master.json"


def norm(text: str) -> str:
    """正規化：去除所有空白，供同源包含比對（字串嚴格相等太脆；此為共同生成欄位的正規化比對）。"""
    return "".join(text.split())


class TestStatsMaster:
    def test_official_counts_match_resume_audit(self):
        d = json.loads(STATS.read_text(encoding="utf-8"))["descriptive"]
        expect = {
            "h7|ALL": (1149, 10952, "10.5%"), "h14|ALL": (1773, 10680, "16.6%"),
            "h7|cabbage": (227, 3357, "6.8%"), "h14|cabbage": (420, 3274, "12.8%"),
            "h7|bok_choy": (443, 3826, "11.6%"), "h14|bok_choy": (682, 3735, "18.3%"),
            "h7|sponge_gourd": (479, 3769, "12.7%"), "h14|sponge_gourd": (671, 3671, "18.3%"),
        }
        for k, (num, den, disp) in expect.items():
            assert d[k]["nonflat"] == num and d[k]["n"] == den and d[k]["percent_display"] == disp, k

    def test_float_boundary_note(self):
        m = json.loads(STATS.read_text(encoding="utf-8"))
        assert m["float_boundary_note"]["h7"]["minus_label"] == 7
        assert m["float_boundary_note"]["h14"]["minus_label"] == 14

    def test_purge_two_segments(self):
        p = json.loads(STATS.read_text(encoding="utf-8"))["purge"]
        assert p["rows"] == 1386 and p["distinct_t"] == 30
        assert p["reasons"]["train_label_not_mature_before_val"] == 695
        assert p["reasons"]["val_label_not_mature_before_test"] == 691


class TestNotesSource:
    def test_19_slides_p01_to_p19(self):
        data = json.loads(NOTES_JSON.read_text(encoding="utf-8"))
        ids = [s["slide_id"] for s in data["slides"]]
        assert ids == [f"p{i:02d}" for i in range(1, 20)]
        assert sum(1 for s in data["slides"] if s["kind"] == "main") == 15

    def test_every_slide_has_script_and_sources(self):
        for s in json.loads(NOTES_JSON.read_text(encoding="utf-8"))["slides"]:
            assert len(s["script"]) >= 30, s["slide_id"]
            assert s.get("sources"), s["slide_id"]

    def test_main_seconds_total(self):
        data = json.loads(NOTES_JSON.read_text(encoding="utf-8"))
        assert sum(s["seconds"] for s in data["slides"] if s["kind"] == "main") == 675


class TestSlideNotesMdSameSource:
    def test_md_contains_all_normalized_scripts(self):
        md = norm((ROOT / "docs/SLIDE_NOTES.md").read_text(encoding="utf-8"))
        for s in json.loads(NOTES_JSON.read_text(encoding="utf-8"))["slides"]:
            assert norm(s["script"]) in md, s["slide_id"]


class TestPptxNotesSameSource:
    def test_pptx_notes_contain_normalized_scripts(self):
        from pptx import Presentation
        prs = Presentation(str(PPTX_V2))
        slides = list(prs.slides)
        assert len(slides) == 19
        data = json.loads(NOTES_JSON.read_text(encoding="utf-8"))
        for s, rec in zip(slides, data["slides"]):
            note = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
            assert norm(rec["script"]) in norm(note), rec["slide_id"]
            for src in rec["sources"]:
                assert norm(src) in norm(note), (rec["slide_id"], src)

    def test_architecture_slide_has_connectors(self):
        from pptx import Presentation
        from pptx.oxml.ns import qn
        prs = Presentation(str(PPTX_V2))
        arch = list(prs.slides)[7]  # p08
        n_cxn = len(arch.shapes._spTree.findall(qn("p:cxnSp")))
        assert n_cxn >= 6, f"架構頁連接箭頭應>=6，實為 {n_cxn}"


class TestDeliverableCounts:
    def test_pdf_pages_equal_19(self):
        from pypdf import PdfReader
        assert len(PdfReader(str(PDF_V2)).pages) == 19

    def test_render_meta_match(self):
        meta = json.loads((RUN2 / "slides_render_v2/render_meta.json").read_text(encoding="utf-8"))
        assert meta["match"] is True and meta["pptx_slides"] == 19


class TestSafetyBehaviors:
    def test_make_assets_refuse_overwrite(self, tmp_path):
        (tmp_path / "x.png").write_bytes(b"1")
        r = subprocess.run([sys.executable, str(ROOT / "slides/source/make_assets.py"),
                            "--out", str(tmp_path), "--stats", str(STATS)],
                           capture_output=True, text=True)
        assert r.returncode == 2 and "REFUSE" in r.stderr

    def test_make_pptx_missing_asset_fails(self, tmp_path):
        empty = tmp_path / "empty_assets"
        empty.mkdir()
        r = subprocess.run([sys.executable, str(ROOT / "slides/source/make_pptx.py"),
                            "--pptx", str(tmp_path / "t.pptx"), "--assets", str(empty),
                            "--notes", str(NOTES_JSON), "--stats", str(STATS)],
                           capture_output=True, text=True)
        assert r.returncode == 3 and "缺圖" in r.stderr

    def test_make_pptx_refuse_overwrite_existing(self):
        r = subprocess.run([sys.executable, str(ROOT / "slides/source/make_pptx.py"),
                            "--pptx", str(PPTX_V2), "--assets", str(ROOT / "slides/assets_v2"),
                            "--notes", str(NOTES_JSON), "--stats", str(STATS)],
                           capture_output=True, text=True)
        assert r.returncode == 2 and "REFUSE" in r.stderr

    def test_v1_outputs_untouched(self):
        import hashlib
        v1 = ROOT / "slides/TW01_菜價預測_研究提案與初步成果_v1.pptx"
        h = hashlib.sha256(v1.read_bytes()).hexdigest()
        rec = json.loads((RUN2 / "protected_before.json").read_text(encoding="utf-8"))
        key = "slides/TW01_菜價預測_研究提案與初步成果_v1.pptx"
        assert rec[key]["sha256"] == h
