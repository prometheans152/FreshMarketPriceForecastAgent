# -*- coding: utf-8 -*-
"""試片專用頁製作：載入 v2 PPTX → 修正第10頁教學標示 → 匯出第1/10頁 PNG。

執行：.venv/Scripts/python.exe tools/video_tools/make_preview_pages.py --run recordings/preview_20261004_01
安全：只在 POWERPNT.EXE 未運行時 Dispatch+Quit；若已運行則 GetActiveObject 且只 Close 簡報、不 Quit 應用。
"""
import argparse
import shutil
import subprocess
import sys
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]
V2 = ROOT / "slides/TW01_菜價預測_研究提案與初步成果_v2.pptx"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--run", required=True)
    args = ap.parse_args()
    run = Path(args.run).resolve()
    run.mkdir(parents=True, exist_ok=True)

    copy = run / "preview_slides.pptx"
    if copy.exists():
        print(f"REFUSE: {copy} 已存在", file=sys.stderr)
        raise SystemExit(2)
    shutil.copy2(V2, copy)

    prs = Presentation(str(copy))
    s10 = prs.slides[9]  # 第10頁 A/B
    changed = {"fixture": 0, "title": 0, "banner": 0}
    for shape in s10.shapes:
        if not shape.has_text_frame:
            continue
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                if "〔封存 fixture 輸出〕" in r.text:
                    r.text = r.text.replace("〔封存 fixture 輸出〕", "〔教學假設〕")
                    changed["fixture"] += 1
                if "數字例（甘藍；教學計算非真店實驗）" in r.text:
                    r.text = r.text.replace("數字例（甘藍；教學計算非真店實驗）", "數字例（甘藍；本頁數值全為教學假設）")
                    changed["title"] += 1
                if "B 實驗性未驗證" in r.text:
                    r.text = r.text + "\n\n※本頁全部數值皆為教學假設，\n非真店實驗，亦非正式模型 fixture 結果"
                    changed["banner"] += 1
    if changed["fixture"] != 2 or changed["title"] != 1 or changed["banner"] != 1:
        print(f"FAIL: 替換計數異常 {changed}（預期 fixture=2,title=1,banner=1）", file=sys.stderr)
        raise SystemExit(3)
    prs.save(str(copy))
    print(f"patched p10: {changed}")

    # --- 安全 COM 匯出第1/10頁 ---
    tl = subprocess.run(["tasklist", "/FI", "IMAGENAME eq POWERPNT.EXE"], capture_output=True)
    already = b"POWERPNT.EXE" in tl.stdout
    import win32com.client
    if already:
        app = win32com.client.GetActiveObject("PowerPoint.Application")
    else:
        app = win32com.client.Dispatch("PowerPoint.Application")
    try:
        pres = app.Presentations.Open(str(copy), WithWindow=False)
        try:
            for idx in (1, 10):
                pres.Slides(idx).Export(str(run / f"preview_slide_{idx:02d}.png"), "PNG", 1600, 900)
        finally:
            pres.Close()  # 只關本簡報
    finally:
        if not already:
            app.Quit()  # 只 Quit 自己啟動的實例
    print("exported preview_slide_01.png / preview_slide_10.png")


if __name__ == "__main__":
    main()
